from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

NAVY      = RGBColor(0x0D, 0x1B, 0x3E)
NAVY_DARK = RGBColor(0x08, 0x10, 0x28)
NAVY_MID  = RGBColor(0x15, 0x28, 0x52)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE    = RGBColor(0xFF, 0x6B, 0x35)
LGRAY     = RGBColor(0xCC, 0xCC, 0xCC)
DIVGRAY   = RGBColor(0x2A, 0x3A, 0x6A)

IMGS = {
    "017": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/e027f003-1000149017.jpg",
    "018": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/e6b79f4a-1000149018.jpg",
    "019": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/8008f10a-1000149019.jpg",
    "020": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/8ad82bc8-1000149020.jpg",
    "021": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/01f972cd-1000149021.jpg",
    "022": "/root/.claude/uploads/e9b5fa60-5927-438b-86d5-288a719c4134/54b6f34c-1000149022.jpg",
}


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def tb(slide, text, l, t, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = "Calibri"
    return box


def tb_multi(slide, lines, l, t, w, h, size, color, bold=False, line_space=None):
    """lines: list of str"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Calibri"


def rect(slide, l, t, w, h, fill_color, line_color=None, line_w=None):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        if line_w:
            shape.line.width = Pt(line_w)
    else:
        shape.line.fill.background()
    return shape


def img(slide, key, l, t, w, h):
    path = IMGS.get(key)
    if path and os.path.exists(path):
        slide.shapes.add_picture(path, l, t, w, h)


def top_bar(slide):
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.07), ORANGE)


def header(slide, label, title):
    top_bar(slide)
    tb(slide, label, Inches(0.6), Inches(0.18), Inches(12), Inches(0.42),
       11, ORANGE, bold=True)
    tb(slide, title, Inches(0.6), Inches(0.55), Inches(12), Inches(0.7),
       24, WHITE, bold=True)
    rect(slide, Inches(0.6), Inches(1.22), Inches(1.4), Inches(0.055), ORANGE)


# ──────────────────────────────────────────
# BUILD
# ──────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# ─── SLIDE 1 — PORTADA ─────────────────────
s = new_slide(prs)
bg(s, NAVY)
rect(s, Inches(10.0), Inches(0), Inches(0.45), Inches(7.5), ORANGE)
rect(s, Inches(10.55), Inches(0), Inches(0.3), Inches(7.5), NAVY_MID)

tb(s, "TRABAJO FINAL — ILUMINACIÓN Y RENDER",
   Inches(0.8), Inches(0.85), Inches(8.8), Inches(0.5), 13, ORANGE, bold=True)

box = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(9), Inches(3))
tf = box.text_frame
tf.word_wrap = True
for i, line in enumerate(["NO SALIR", "DE DÍA"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = line
    r.font.size = Pt(64)
    r.font.color.rgb = WHITE
    r.font.bold = True
    r.font.name = "Calibri"

tb(s, "Escenario de iluminación cinematográfica  ·  Unreal Engine 5",
   Inches(0.8), Inches(4.35), Inches(8.8), Inches(0.55), 18, LGRAY)
tb(s, "Modalidad A — Opción Cinematográfica  ·  4 secuencias",
   Inches(0.8), Inches(4.95), Inches(8.8), Inches(0.45), 14, ORANGE)
rect(s, Inches(0.8), Inches(5.55), Inches(3), Inches(0.05), LGRAY)
tb(s, "EUNEIZ  ·  Grado en Arte para Videojuegos  ·  Curso 2025–26",
   Inches(0.8), Inches(5.75), Inches(8.8), Inches(0.5), 12, LGRAY)

# ─── SLIDE 2 — CONCEPTO ────────────────────
s = new_slide(prs)
bg(s, NAVY)
header(s, "CONCEPTO Y NARRATIVA", "\"En combate, la luz del día es el enemigo\"")

tb(s,
   "Cuatro secuencias cinematográficas narran la espera en una trinchera de la Primera Guerra Mundial. "
   "Salir durante el día significa muerte segura: la luz, normalmente asociada a la vida, se convierte aquí "
   "en amenaza. La oscuridad del atardecer es el único momento de acción posible.",
   Inches(0.6), Inches(1.45), Inches(5.9), Inches(2.0), 13, LGRAY)

seqs = [
    ("SEC. 1", "ESTABLISHING", "Plano general.\nContextualiza el espacio."),
    ("SEC. 2", "ARSENAL",      "Close-up DoF.\nImplica presencia humana."),
    ("SEC. 3", "TRAVELLING",   "Dolly + cartel.\n\"Keep to the trench.\""),
    ("SEC. 4", "SUNSET",       "Exterior. Sol en horizonte.\nLa noche llega."),
]
for i, (num, title, desc) in enumerate(seqs):
    x = Inches(0.6 + i * 3.15)
    y = Inches(3.5)
    rect(s, x, y, Inches(2.9), Inches(3.4), NAVY_MID, ORANGE, 1)
    tb(s, num,   x+Inches(0.18), y+Inches(0.18), Inches(1), Inches(0.38), 10, ORANGE, bold=True)
    tb(s, title, x+Inches(0.18), y+Inches(0.58), Inches(2.6), Inches(0.5), 15, WHITE, bold=True)
    rect(s, x+Inches(0.18), y+Inches(1.12), Inches(2.6), Inches(0.04), ORANGE)
    tb(s, desc,  x+Inches(0.18), y+Inches(1.25), Inches(2.6), Inches(1.6), 11, LGRAY)

# ─── SLIDE 3 — SETUP TÉCNICO ───────────────
s = new_slide(prs)
bg(s, NAVY)
header(s, "SETUP TÉCNICO", "Motor · Luces · Render")

left = [
    ("MOTOR",             "Unreal Engine 5"),
    ("ILUMINACIÓN GLOBAL", "Lumen — Dynamic GI + Reflections"),
    ("SOMBRAS",           "Virtual Shadow Maps"),
    ("PIPELINE DE RENDER","Movie Render Queue (offline)"),
    ("FRAME RATE",        "24 fps — look cinematográfico"),
    ("ANTI-ALIASING",     "TSR — Temporal Super Resolution"),
    ("ASSET BASE",        "FAB Trench Pack — luces / cámaras\nrehechas íntegramente"),
]
right = [
    ("KEY LIGHT",         "Directional Light — Movable"),
    ("FILL LIGHT",        "Sky Light — Real Time Capture"),
    ("COLOR TEMPERATURE", "3.500 K  (ámbar cálido)"),
    ("VOLUMETRIC FOG",    "Exponential Height Fog  ·  Density 0.02"),
    ("LENTES",            "35 mm  ·  50 mm  (Cine Camera Actor)"),
    ("APERTURA",          "f/2  →  f/8  según secuencia"),
    ("RESOLUCIÓN",        "1920 × 1080 — ratio 16:9"),
]

for i, (lbl, val) in enumerate(left):
    y = Inches(1.45) + Inches(i * 0.84)
    tb(s, lbl, Inches(0.6), y, Inches(5.5), Inches(0.3), 9, ORANGE, bold=True)
    tb(s, val, Inches(0.6), y+Inches(0.3), Inches(5.5), Inches(0.48), 12, WHITE)

for i, (lbl, val) in enumerate(right):
    y = Inches(1.45) + Inches(i * 0.84)
    tb(s, lbl, Inches(7.0), y, Inches(5.8), Inches(0.3), 9, ORANGE, bold=True)
    tb(s, val, Inches(7.0), y+Inches(0.3), Inches(5.8), Inches(0.48), 12, WHITE)

rect(s, Inches(6.55), Inches(1.3), Inches(0.05), Inches(6.0), DIVGRAY)

# ─── SLIDE 4 — ILUMINACIÓN ─────────────────
s = new_slide(prs)
bg(s, NAVY)
header(s, "ILUMINACIÓN", "Key  ·  Fill  ·  Atmósfera volumétrica")

lights = [
    ("KEY LIGHT", "Directional Light",
     ["Mobility: Movable",
      "Intensity: 8–10 lux",
      "Color Temp: 3.500 K",
      "Source Angle: 1.5",
      "Atmosphere Sun Light: ON",
      "Cast Volumetric Shadow: ON",
      "Vol. Scattering Int.: 2.0"]),
    ("FILL / SKY", "Sky Light",
     ["Source: Real Time Capture",
      "Intensity Scale: 0.4",
      "Sky Atmosphere vinculada",
      "No compite con la key",
      "Evita negros absolutos",
      "Lower Hemisphere: Color"]),
    ("ATMÓSFERA", "Volumetric Fog + Lumen",
     ["Fog Density: 0.02",
      "Scattering Dist.: 0.6",
      "Inscattering: ámbar oscuro",
      "Lumen GI Quality: 2",
      "God rays vía Vol. Shadow",
      "Niebla en zonas de haces"]),
]
for i, (title, sub, items) in enumerate(lights):
    x = Inches(0.55 + i * 4.25)
    y = Inches(1.45)
    rect(s, x, y, Inches(4.0), Inches(5.7), NAVY_MID, ORANGE, 1)
    tb(s, title, x+Inches(0.2), y+Inches(0.18), Inches(3.6), Inches(0.42), 14, ORANGE, bold=True)
    tb(s, sub,   x+Inches(0.2), y+Inches(0.62), Inches(3.6), Inches(0.38), 12, WHITE)
    rect(s, x+Inches(0.2), y+Inches(1.08), Inches(3.6), Inches(0.04), ORANGE)
    for j, item in enumerate(items):
        tb(s, "· "+item, x+Inches(0.2), y+Inches(1.22)+Inches(j*0.63),
           Inches(3.6), Inches(0.58), 11, LGRAY)

# ─── SLIDE 5 — SECUENCIA 1 ────────────────
s = new_slide(prs)
bg(s, NAVY_DARK)
header(s, "SECUENCIA 1 — ESTABLISHING SHOT", "Plano general · Presentación del escenario")

img(s, "017", Inches(0.3), Inches(1.45), Inches(7.9), Inches(4.45))

data = [
    ("TIPO DE PLANO",  "Plano General (Establishing Shot)"),
    ("LENTE",          "35 mm — ángulo amplio"),
    ("APERTURA",       "f/8 — profundidad de campo total"),
    ("COMPOSICIÓN",    "Líneas guía convergentes · Punto\nde fuga central (pasillo trinchera)"),
    ("KEY LIGHT",      "Directional lateral-superior\nGod rays visibles por Vol. Fog"),
    ("COLOR",          "Dominante ámbar 3500 K\nSky fill en sombras neutras"),
    ("RATIO",          "2.39:1 Cinemascope (letterbox)"),
]
y0 = Inches(1.45)
for lbl, val in data:
    tb(s, lbl, Inches(8.55), y0, Inches(4.4), Inches(0.28), 9, ORANGE, bold=True)
    tb(s, val, Inches(8.55), y0+Inches(0.28), Inches(4.4), Inches(0.5), 11, WHITE)
    y0 += Inches(0.84)

# ─── SLIDE 6 — SECUENCIA 2 ────────────────
s = new_slide(prs)
bg(s, NAVY_DARK)
header(s, "SECUENCIA 2 — ARSENAL", "Close-up · Cajas de munición · DoF extremo")

img(s, "019", Inches(0.3), Inches(1.45), Inches(7.9), Inches(4.45))

data = [
    ("TIPO DE PLANO",  "Primer plano con bajo ángulo\n(cámara a nivel del suelo)"),
    ("LENTE",          "50 mm — compresión de perspectiva"),
    ("APERTURA",       "f/2 — bokeh pronunciado"),
    ("FOCO",           "Manual focus en cajas\nForeground desenfocado"),
    ("ILUMINACIÓN",    "Rim light cálido lateral\nFondo en penumbra (clave baja)"),
    ("NARRATIVA",      "Escala humana implícita\nVida sin mostrar personas"),
    ("COLOR",          "Ámbar saturado sobre negro\nContraste térmico máximo"),
]
y0 = Inches(1.45)
for lbl, val in data:
    tb(s, lbl, Inches(8.55), y0, Inches(4.4), Inches(0.28), 9, ORANGE, bold=True)
    tb(s, val, Inches(8.55), y0+Inches(0.28), Inches(4.4), Inches(0.5), 11, WHITE)
    y0 += Inches(0.84)

# ─── SLIDE 7 — SECUENCIA 3 ────────────────
s = new_slide(prs)
bg(s, NAVY_DARK)
header(s, "SECUENCIA 3 — TRAVELLING + CARTEL", "Dolly forward · Pull focus · \"Keep to the trench in daylight\"")

img(s, "020", Inches(0.3),  Inches(1.5), Inches(6.1), Inches(3.45))
img(s, "021", Inches(6.65), Inches(1.5), Inches(6.3), Inches(3.45))

tb(s, "Inicio — dolly con god rays a contraluz",
   Inches(0.3), Inches(5.05), Inches(6.1), Inches(0.38), 10, LGRAY)
tb(s, "Final — close-up cartel en penumbra total",
   Inches(6.65), Inches(5.05), Inches(6.3), Inches(0.38), 10, LGRAY)

notes = ("Lente: 35 mm → 50 mm  ·  Apertura: f/4 → f/2  ·  "
         "Movimiento: Dolly forward (Sequencer keyframes, ease-in/out)  ·  "
         "Iluminación: contraluz key + Vol. Fog al inicio → oscuridad total al llegar al cartel  ·  "
         "Pull focus dinámico: de la niebla al cartel")
tb(s, notes, Inches(0.3), Inches(5.6), Inches(12.7), Inches(0.75), 10, ORANGE)

# ─── SLIDE 8 — SECUENCIA 4 ────────────────
s = new_slide(prs)
bg(s, NAVY_DARK)
header(s, "SECUENCIA 4 — SUNSET", "Plano exterior · El sol se pone · La noche llega")

img(s, "022", Inches(0.3), Inches(1.45), Inches(7.9), Inches(4.45))

data = [
    ("TIPO DE PLANO",  "Plano general exterior\nVista desde dentro de la trinchera"),
    ("LENTE",          "35 mm — campo amplio"),
    ("APERTURA",       "f/8 — todo en foco"),
    ("ILUMINACIÓN",    "Directional casi en horizonte\nSol a ~2–3° — máxima calidez"),
    ("COLOR",          "Rojo-naranja dominante\nSiluetas en negro absoluto"),
    ("VOLUMETRIC",     "Humo + fog denso\nAtmósfera de campo de batalla"),
    ("NARRATIVA",      "La oscuridad = libertad de acción\nEl sol cayendo = cuenta atrás"),
]
y0 = Inches(1.45)
for lbl, val in data:
    tb(s, lbl, Inches(8.55), y0, Inches(4.4), Inches(0.28), 9, ORANGE, bold=True)
    tb(s, val, Inches(8.55), y0+Inches(0.28), Inches(4.4), Inches(0.5), 11, WHITE)
    y0 += Inches(0.84)

# ─── SLIDE 9 — PALETA Y MOOD ──────────────
s = new_slide(prs)
bg(s, NAVY)
header(s, "PALETA CROMÁTICA Y MOOD", "Psicología del color aplicada")

swatches = [
    (RGBColor(0xD4,0x7A,0x00), "ÁMBAR / OCRE",
     "Dominante cálida\n3500K Directional\nPeligro · Urgencia · Calor"),
    (RGBColor(0xE8,0x40,0x10), "ROJO-NARANJA",
     "Sunset (sec. 4)\nSol en horizonte\nSangre · Guerra · Alerta"),
    (RGBColor(0x15,0x10,0x08), "NEGRO PROFUNDO",
     "Sombras Lumen\nSiluetas / trinchera\nMuerte · Opresión"),
    (RGBColor(0x8A,0x70,0x40), "TIERRA / SEPIA",
     "Base color del asset\nSacos · Madera · Arena\nGuerra · Abandono"),
]
for i, (color, name, desc) in enumerate(swatches):
    x = Inches(0.55 + i * 3.2)
    rect(s, x, Inches(1.55), Inches(3.0), Inches(1.7), color)
    tb(s, name, x, Inches(3.35), Inches(3.0), Inches(0.42), 11, WHITE, bold=True)
    tb(s, desc, x, Inches(3.82), Inches(3.0), Inches(1.3), 10, LGRAY)

rect(s, Inches(0.55), Inches(5.45), Inches(12.2), Inches(0.05), DIVGRAY)
mood = ("La inversión psicológica del color es el eje narrativo: "
        "la LUZ cálida (ámbar, rojo) representa el PELIGRO, "
        "mientras la OSCURIDAD es la única vía de supervivencia. "
        "Lumen amplifica este contraste con rebotes que saturan las zonas iluminadas "
        "y hunden las sombras en negro profundo.")
tb(s, mood, Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.6), 12, LGRAY)

# ─── SLIDE 10 — CRITERIOS ──────────────────
s = new_slide(prs)
bg(s, NAVY)
header(s, "CRITERIOS DE EVALUACIÓN", "Alineación del trabajo con los criterios de la asignatura")

criteria = [
    (40, "ILUMINACIÓN",
     "Directional Light 3500K + Sky Light Real Time Capture + Lumen GI. "
     "Volumetric Fog (density 0.02) con Cast Volumetric Shadow para god rays. "
     "Temperatura y ángulo solar consistentes en las 4 secuencias."),
    (20, "MOOD Y ATMÓSFERA",
     "Paleta ámbar/rojo/negro. Niebla volumétrica que refuerza la tensión. "
     "Inversión psicológica: luz = peligro, oscuridad = escape. "
     "Presencia humana implícita sin figuras."),
    (15, "COMPOSICIÓN",
     "Líneas guía (paredes de trinchera). Regla de tercios. Variedad de encuadres: "
     "general → close-up → dolly → exterior. Ratio 2.39:1 con letterbox."),
    (10, "ASPECTOS TÉCNICOS",
     "Lumen + Virtual Shadow Maps. Movie Render Queue 24fps. "
     "Lentes 35mm/50mm. Apertura f/2–f/8. Asset FAB con iluminación reconstruida íntegramente."),
    (15, "SEGUIMIENTO",
     "4 secuencias completas entregadas. Breakdown del motor. "
     "Narrativa coherente y documentada desde el concepto hasta el render final."),
]

y0 = Inches(1.45)
for pct, title, desc in criteria:
    badge = rect(s, Inches(0.5), y0, Inches(0.88), Inches(0.7), ORANGE)
    tb(s, f"{pct}%", Inches(0.5), y0+Inches(0.1), Inches(0.88), Inches(0.5),
       18, WHITE, bold=True, align=PP_ALIGN.CENTER)
    tb(s, title, Inches(1.55), y0,            Inches(2.5),  Inches(0.38), 12, ORANGE, bold=True)
    tb(s, desc,  Inches(1.55), y0+Inches(0.38), Inches(11.2), Inches(0.5),  10, LGRAY)
    y0 += Inches(1.05)

# ─── SAVE ──────────────────────────────────
out = "/home/user/IAEntorno/NO_SALIR_DE_DIA_presentacion.pptx"
prs.save(out)
print("OK:", out)
